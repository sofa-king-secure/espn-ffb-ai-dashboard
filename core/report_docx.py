"""Word (.docx) export via python-docx. Returns bytes for st.download_button."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

from .dossier import spread_text
from .espn_client import NON_STARTING_SLOTS

INK = RGBColor(0x1B, 0x1B, 0x1B)
SIGNAL = RGBColor(0xC2, 0x41, 0x0C)


def _shade(cell, hex_fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(str(h))
        run.bold = True
        run.font.size = Pt(9)
        _shade(c, "E7E2D6")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t


def _inline(paragraph, text):
    """Render **bold** segments inside a paragraph."""
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part.replace("`", ""))


def _markdown(doc, md: str):
    table_buf = []

    def flush_table():
        if not table_buf:
            return
        rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in table_buf
                if not re.match(r"^\|?\s*:?-{2,}", ln.strip())]
        if rows:
            _table(doc, rows[0], rows[1:])
        table_buf.clear()

    for line in md.splitlines():
        s = line.rstrip()
        if s.strip().startswith("|"):
            table_buf.append(s)
            continue
        flush_table()
        if not s.strip():
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            doc.add_heading(m.group(2).replace("**", ""), level=min(len(m.group(1)) + 1, 4))
        elif re.match(r"^\s*[-*]\s+", s):
            level = (len(s) - len(s.lstrip())) // 2
            p = doc.add_paragraph(style="List Bullet 2" if level else "List Bullet")
            _inline(p, re.sub(r"^\s*[-*]\s+", "", s))
        elif re.match(r"^\s*\d+[.)]\s+", s):
            p = doc.add_paragraph(style="List Number")
            _inline(p, re.sub(r"^\s*\d+[.)]\s+", "", s))
        else:
            _inline(doc.add_paragraph(), s)
    flush_table()


def build_docx(snap: dict, ai_report: str = "", ai_model: str = "", lineup_moves: list | None = None) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.8))
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10.5)

    m, mu = snap["meta"], snap["matchup"]
    title = doc.add_heading(f"{m['team_name']} — Week {m['week']} Game Plan", level=0)
    title.runs[0].font.color.rgb = INK
    doc.add_paragraph(f"{m.get('league_name') or 'League ' + str(m['league_id'])} · Record {m['record']} · "
                      f"Snapshot {m['fetched_at']}")

    doc.add_heading("Matchup", level=1)
    _table(doc, ["", m["team_name"], mu["opponent"]],
           [["Projected", mu["my_projected"], mu["opp_projected"]],
            ["Live score", mu["my_score"], mu["opp_score"]]], widths=[1.3, 2.6, 2.6])
    p = doc.add_paragraph()
    r = p.add_run(spread_text(mu["spread"]))
    r.bold = True
    r.font.color.rgb = SIGNAL if mu["spread"] < 0 else INK

    if snap["alerts"]:
        doc.add_heading("Action required", level=1)
        for a in snap["alerts"]:
            doc.add_paragraph(f"[{a['level'].upper()}] {a['text']}", style="List Bullet")

    if snap["pending"]:
        doc.add_heading("Pending moves", level=1)
        _table(doc, ["Type", "Status", "Processes", "Partner", "In", "Out"],
               [[x["type"], x["status"], x["process_date"], x["partner"],
                 ", ".join(x["incoming"]) or "—", ", ".join(x["outgoing"]) or "—"] for x in snap["pending"]])

    hdr = ["Slot", "Player", "Pos", "Opp", "Status", "Proj", "Pts"]
    for heading, grp in (("Starting lineup", [x for x in snap["roster"] if x["slot_id"] not in NON_STARTING_SLOTS]),
                         ("Bench & IR", [x for x in snap["roster"] if x["slot_id"] in NON_STARTING_SLOTS])):
        doc.add_heading(heading, level=1)
        _table(doc, hdr, [[x["slot"], x["name"] + (" (locked)" if x["locked"] else ""), x["pos"], x["opponent"],
                           x["status"] or "OK", x["projected"], x["actual"]] for x in grp],
               widths=[0.6, 2.2, 0.5, 0.7, 0.7, 0.6, 0.6])

    if lineup_moves:
        doc.add_heading("Recommended lineup changes", level=1)
        _table(doc, ["Player", "From", "To", "Proj"],
               [[x["name"], x["from"], x["to"], x["projected"]] for x in lineup_moves])

    movers = sorted([f for f in snap["free_agents"] if f.get("pct_change") is not None],
                    key=lambda f: f["pct_change"], reverse=True)[:12]
    if movers:
        doc.add_heading("Waiver wire — top % add", level=1)
        _table(doc, ["Player", "Pos", "Team", "Proj", "% Rost", "% Chg"],
               [[f["name"], f["pos"], f["pro_team"], f["projected"], f"{f['pct_owned']}%",
                 f"{f['pct_change']:+.2f}%"] for f in movers])

    if ai_report:
        doc.add_page_break()
        doc.add_heading(f"AI strategy report{f' ({ai_model})' if ai_model else ''}", level=1)
        _markdown(doc, ai_report)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
